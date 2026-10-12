"""Yaadash Library — my study notes as an offline, cited answer source (SIH-4 / Y5).

THE PROBLEM IT SOLVES
Smart India Hackathon 2025 SIH25231 (NTRO) asks for an offline RAG over documents
that answers with numbered citations. For an engineering student that is the
semester: lecture PDFs, slides, a lab manual in Word, Markdown notes. "What do my
notes say about setup time?" should come back with an answer AND the page it came
from, so I can check it — and with "I don't know" when the notes do not cover it,
because a 4B model asked about something outside them will happily make it up.

THE PIPELINE
    file -> extract (per page / slide / heading) -> chunk (~220 words, overlap)
         -> embed (bge-small, the same CPU embedder as yaadash_v2) -> LanceDB table
    question -> BM25 + vector -> Reciprocal Rank Fusion -> abstention gate
             -> numbered passages -> local LLM (injectable) -> answer + [n] -> file, page

Retrieval mirrors yaadash_recall.py: BM25 for rare exact tokens ("74LS00", "pin 14"),
vectors for paraphrase, fused with RRF (1/(K + rank)) so neither score needs
calibrating against the other, and a gate that returns nothing rather than the
least-bad match.

WHAT IT DOES NOT DO (said here so nobody assumes it)
- Scanned pages (no text layer) are listed as scanned. OCR runs only when PyMuPDF
  finds a local Tesseract; this laptop has none, so they are not read.
- Word files have no fixed pages; a DOCX citation names its heading, plus the page
  Word last laid it out on when the file records that (Word-saved files do).
- Images, diagrams and equations rendered as pictures are not read.

Storage lives in its own folder (default data/yaadash/library, git-ignored;
JARVIS_LIBRARY_DIR or the constructor moves it): lancedb/ holds the passages,
manifest.json the documents. Nothing here touches the memories table.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import sys
import threading
import time
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Callable, Optional, Protocol

import numpy as np

# ------------------------------------------------------------------ configuration
SUPPORTED = {".pdf": "pdf", ".docx": "docx", ".pptx": "pptx", ".md": "md",
             ".markdown": "md", ".txt": "txt"}
TABLE = "library_chunks"
MANIFEST = "manifest.json"

# Chunks: ~200-250 words. Long enough to hold a definition with its example, short
# enough that five of them (~1.7k tokens) fit the 4B model's prompt with room to
# answer. Overlap carries the last sentence or two across a boundary so a fact
# split between chunks is whole in at least one of them. Chunks never cross a page:
# a citation has to name ONE page.
TARGET_WORDS = 220
MAX_WORDS = 250
OVERLAP_WORDS = 40
MIN_TAIL_WORDS = 60          # a shorter last piece is folded into the chunk before it

MAX_FILE_BYTES = 200 * 1024 * 1024
MAX_FILES = 2000             # one ingest of a folder; a whole drive is not a library
SKIP_DIRS = {".git", ".obsidian", ".trash", "node_modules", "__pycache__", ".venv",
             "venv", "$recycle.bin", "system volume information", "appdata"}

RRF_K = 60                   # the standard RRF constant, as in yaadash_recall
CANDIDATES = 30              # hits taken from each system before fusion
DEFAULT_MODEL = "qwen3:4b-instruct-2507-q4_K_M"   # the brain's model (core/agent.py)
OLLAMA = "http://127.0.0.1:11434"

NOT_IN_NOTES = "I don't know — your notes don't cover that."


def default_root() -> Path:
    """data/yaadash/library unless JARVIS_LIBRARY_DIR says otherwise."""
    env = os.environ.get("JARVIS_LIBRARY_DIR")
    if env:
        return Path(env)
    from home import YAADASH_DATA          # core/home.py: data/yaadash, git-ignored
    return Path(YAADASH_DATA) / "library"


# ------------------------------------------------------------------ text helpers
_TOKEN = re.compile(r"[a-z0-9]+")
_WS = re.compile(r"\s+")


def _stem(w: str) -> str:
    """The same light suffix stripping as yaadash_recall._stem ("complete" meets
    "completed", "registers" meets "register"), kept symmetric between corpus and
    query. Not a full stemmer: Porter mangles Hinglish and part numbers."""
    if len(w) > 5:
        for suf in ("ing", "ed", "es"):
            if w.endswith(suf):
                w = w[: -len(suf)]
                break
        else:
            if w.endswith("s") and not w.endswith("ss"):
                w = w[:-1]
    if len(w) > 5 and w.endswith("e"):
        w = w[:-1]
    return w


def tokens(text: str) -> list[str]:
    """Lower-case word tokens for BM25. Single letters are kept on purpose: in
    lecture notes "D", "T", "JK" and "SR" name flip-flops. Contractions are undone
    first so "don't" does not leave a stray "t" that looks like a T flip-flop."""
    t = (text or "").lower().replace("’", "'")
    t = re.sub(r"n't\b", " not", t)
    t = re.sub(r"'s\b", "", t)
    return [_stem(w) for w in _TOKEN.findall(t)]


# Words that carry no search signal in a QUESTION (the corpus keeps them, so IDF
# stays honest). English question scaffolding plus the Hinglish I actually say.
# "state", "gate", "time" and "clock" are NOT here: in these notes they are topics.
_QUERY_STOP = set(tokens("""
a an the is are was were be been being am of in on at to for from by with and or but
not no what which who whom whose when where why how does do did done can could should
would will shall may might must explain describe define definition meaning mean
tell me my mine i you your we our us it its this that these those there their they them
he she his her about between difference vs versus give list name write short note notes
according lecture lectures slide slides say says said kind type please pls jarvis
kya hai hain ka ki ke ko se mein main aur ya bhi toh kaise kyun kyon kab kahan kaun
batao bata samjhao samjha mera meri mere mujhe bhai yaar hota hoti hote matlab wala
wali wale karte karta karna kare kr krna na ji tha thi the kitna kitni kitne konsa
kaunsa konsi kaunsi farak fark antar tak liye
"""))


def query_terms(text: str) -> list[str]:
    q = tokens(text)
    return [t for t in q if t not in _QUERY_STOP] or q


def _clean_ws(text: str) -> str:
    return _WS.sub(" ", text or "").strip()


def _word_count(text: str) -> int:
    return len((text or "").split())


# Sentence ends: punctuation followed by a capital, digit or bracket; bullets too.
_SENT_END = re.compile(r"(?<=[.!?])\s+(?=[\"'(\[A-Z0-9])|\s*[•▪●]\s+")


def _sentences(text: str) -> list[str]:
    return [s.strip() for s in _SENT_END.split(_clean_ws(text)) if s and s.strip()]


def chunk_text(text: str, target: int = TARGET_WORDS, max_words: int = MAX_WORDS,
               overlap: int = OVERLAP_WORDS, min_tail: int = MIN_TAIL_WORDS) -> list[str]:
    """Pack whole sentences into ~target-word chunks; carry <= `overlap` words of
    trailing sentences into the next chunk. A sentence longer than max_words (a
    table flattened into one line) is cut by words."""
    pieces: list[str] = []
    for s in _sentences(text):
        w = s.split()
        if len(w) <= max_words:
            pieces.append(s)
        else:
            pieces += [" ".join(w[i:i + max_words]) for i in range(0, len(w), max_words)]

    def words(lst: list[str]) -> int:
        return sum(_word_count(s) for s in lst)

    chunks: list[list[str]] = []
    cur: list[str] = []
    n_carry = 0                     # leading sentences of `cur` repeated from the last chunk

    def close() -> None:
        nonlocal cur, n_carry
        chunks.append(cur)
        carry: list[str] = []
        for t in reversed(cur):
            if words(carry) + _word_count(t) > overlap:
                break
            carry.insert(0, t)
        if len(carry) == len(cur):  # never repeat a whole chunk
            carry = []
        cur, n_carry = list(carry), len(carry)

    for s in pieces:
        n = _word_count(s)
        if words(cur) + n > max_words:
            if len(cur) > n_carry:
                close()
            if words(cur) + n > max_words:          # the overlap itself does not fit
                cur, n_carry = [], 0
        cur.append(s)
        if words(cur) >= target:
            close()
    fresh = cur[n_carry:]
    if fresh:
        if chunks and words(fresh) < min_tail:
            # Fold a short tail into the chunk before it rather than index a
            # 20-word scrap on its own (its overlap is already in that chunk).
            chunks[-1] = chunks[-1] + fresh
        else:
            chunks.append(cur)
    return [" ".join(c) for c in chunks if c]


# ------------------------------------------------------------------ extraction
@dataclass
class Section:
    page: int                # 1-based page / slide; 0 when the format has no pages
    loc: str                 # what a citation calls it: "p. 4", "slide 3", "§ Flip-flops"
    heading: str
    text: str


@dataclass
class Extracted:
    kind: str
    title: str
    pages: int = 0           # pages or slides; 0 for formats without pages
    sections: list[Section] = field(default_factory=list)
    scanned: list[int] = field(default_factory=list)    # no text layer, has images
    blank: list[int] = field(default_factory=list)
    ocr: list[int] = field(default_factory=list)        # text came from OCR (estimated)
    notes: list[str] = field(default_factory=list)


_ocr_state: dict = {}


def ocr_available() -> bool:
    """True only when PyMuPDF can find a local Tesseract (never downloads one)."""
    if "ok" not in _ocr_state:
        try:
            import fitz
            _ocr_state["ok"] = bool(fitz.get_tessdata())
        except Exception:
            _ocr_state["ok"] = False
    return _ocr_state["ok"]


def _ocr_page(page) -> str:
    tp = page.get_textpage_ocr(language="eng", dpi=200, full=True)
    return page.get_text("text", textpage=tp)


MIN_PAGE_WORDS = 3            # fewer extractable words than this: no text layer


def _extract_pdf(path: Path) -> Extracted:
    import fitz
    with fitz.open(str(path)) as doc:
        title = (doc.metadata or {}).get("title") or path.stem
        out = Extracted("pdf", title=title.strip() or path.stem, pages=doc.page_count)
        flags = fitz.TEXTFLAGS_TEXT | fitz.TEXT_DEHYPHENATE
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text", flags=flags)
            if _word_count(text) < MIN_PAGE_WORDS:
                if not page.get_images(full=False):
                    out.blank.append(i)
                    continue
                if ocr_available():
                    try:
                        text = _ocr_page(page)
                    except Exception as e:           # a bad page must not lose the file
                        out.notes.append(f"p. {i}: OCR failed ({type(e).__name__})")
                        text = ""
                    if _word_count(text) >= MIN_PAGE_WORDS:
                        out.ocr.append(i)
                    else:
                        out.scanned.append(i)
                        continue
                else:
                    out.scanned.append(i)
                    continue
            out.sections.append(Section(i, f"p. {i}", "", _clean_ws(text)))
    if out.scanned:
        out.notes.append(f"{len(out.scanned)} scanned page(s) not read: OCR not available "
                         f"(no local Tesseract)")
    return out


def _extract_pptx(path: Path) -> Extracted:
    from pptx import Presentation
    prs = Presentation(str(path))
    slides = list(prs.slides)
    out = Extracted("pptx", title=(prs.core_properties.title or path.stem).strip() or path.stem,
                    pages=len(slides))
    for i, slide in enumerate(slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = _clean_ws(slide.shapes.title.text_frame.text)
        except Exception:
            title = ""
        parts: list[str] = []
        for shp in slide.shapes:
            if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
                t = _clean_ws(shp.text_frame.text)
                if t and t != title:
                    parts.append(t if t.endswith((".", "!", "?", ":")) else t + ".")
            if getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    cells = [_clean_ws(c.text) for c in row.cells]
                    parts.append(" | ".join(c for c in cells if c) + ".")
        if slide.has_notes_slide:
            notes = _clean_ws(slide.notes_slide.notes_text_frame.text)
            if notes:
                parts.append("Speaker notes: " + notes)
        body = " ".join(parts)
        if _word_count(title + " " + body) < MIN_PAGE_WORDS:
            (out.scanned if any(s.shape_type == 13 for s in slide.shapes) else out.blank).append(i)
            continue
        text = (title + ". " + body) if title else body
        out.sections.append(Section(i, f"slide {i}", title, text))
    if out.scanned:
        out.notes.append(f"{len(out.scanned)} picture-only slide(s) not read")
    return out


def _heading_label(trail: list[str]) -> str:
    return " › ".join(trail[-2:])


def _extract_docx(path: Path) -> Extracted:
    """Paragraphs and tables in document order, sectioned by Heading/Title styles.

    Pages: a .docx has no fixed pages, but Word records where it last broke them
    (w:lastRenderedPageBreak) and authors insert hard breaks. When either exists the
    citation adds that page; otherwise it names the heading alone.
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))
    title = (d.core_properties.title or "").strip() or path.stem
    out = Extracted("docx", title=title)
    body = d.element.body
    rendered = bool(body.xpath(".//w:lastRenderedPageBreak"))
    breaks_q = ".//w:lastRenderedPageBreak" if rendered else ".//w:br[@w:type='page']"
    has_pages = bool(body.xpath(breaks_q))
    page = 1
    trail: list[str] = []
    buf: list[str] = []
    buf_page = 1

    def flush():
        text = " ".join(buf).strip()
        if _word_count(text) >= 1:
            label = _heading_label(trail)
            loc = f"§ {label}" if label else "§ (start)"
            if has_pages:
                loc += f" · p. {buf_page}"
            out.sections.append(Section(buf_page if has_pages else 0, loc, label, text))
        buf.clear()

    for block in d.iter_inner_content():
        if isinstance(block, Paragraph):
            # A paragraph that holds a break is credited to the page after it: a
            # hard break sits in its own empty paragraph, and Word's rendered break
            # usually opens the first paragraph of the new page.
            page += len(block._p.xpath(breaks_q))
            text = _clean_ws(block.text)
            style = (block.style.name if block.style is not None else "") or ""
            level = None
            if style == "Title":
                level = 0
            elif style.startswith("Heading"):
                m = re.search(r"(\d+)", style)
                level = int(m.group(1)) if m else 1
            if level is not None and text:
                flush()
                trail = trail[:max(level - 1, 0)] if level else []
                trail.append(text)
                continue
            if text:
                if not buf:
                    buf_page = page
                buf.append(text if text.endswith((".", "!", "?", ":")) else text + ".")
        elif isinstance(block, Table):
            if not buf:
                buf_page = page
            for row in block.rows:
                cells = []
                for c in row.cells:
                    t = _clean_ws(c.text)
                    if t and (not cells or cells[-1] != t):      # merged cells repeat
                        cells.append(t)
                if cells:
                    buf.append(" | ".join(cells) + ".")
    flush()
    out.pages = page if has_pages else 0
    return out


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")


def _md_inline(text: str) -> str:
    """Markdown and Obsidian syntax -> plain words (the text is for reading aloud
    and for BM25, not for rendering)."""
    t = re.sub(r"!\[[^\]]*\]\([^)]*\)", " ", text)                  # images
    t = re.sub(r"!\[\[[^\]]*\]\]", " ", t)                           # embeds
    t = re.sub(r"\[\[([^\]|]+)\|([^\]]+)\]\]", r"\2", t)             # [[note|alias]]
    t = re.sub(r"\[\[([^\]]+)\]\]", r"\1", t)                         # [[note]]
    t = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", t)                    # [text](url)
    t = re.sub(r"(\*\*|__|`)", "", t)
    t = re.sub(r"^\s*(?:[-*+]|\d+[.)])\s+", "", t)                    # list markers
    t = re.sub(r"^\s*>\s?", "", t)                                    # quotes
    return t


def _extract_md(path: Path) -> Extracted:
    text = _read_text_file(path)
    title = path.stem
    fm = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if fm:
        m = re.search(r"^title:\s*['\"]?(.+?)['\"]?\s*$", fm.group(1), re.M)
        if m:
            title = m.group(1)
        text = text[fm.end():]
    out = Extracted("md", title=title)
    trail: list[str] = []
    buf: list[str] = []
    fenced = False

    def flush():
        body = " ".join(buf).strip()
        if _word_count(body) >= 1:
            label = _heading_label(trail)
            out.sections.append(Section(0, f"§ {label}" if label else "§ (start)", label, _clean_ws(body)))
        buf.clear()

    for line in text.splitlines():
        if _MD_FENCE.match(line):
            fenced = not fenced
            continue
        m = None if fenced else _MD_HEADING.match(line)
        if m:
            flush()
            level = len(m.group(1))
            trail = trail[:level - 1]
            trail.append(_clean_ws(_md_inline(m.group(2))))
            continue
        ln = _md_inline(line).strip()
        if ln:
            buf.append(ln if fenced or ln.endswith((".", "!", "?", ":", "|")) else ln + ".")
        elif buf and not fenced:
            buf.append("")
    flush()
    return out


TXT_SECTION_WORDS = MAX_WORDS


def _extract_txt(path: Path) -> Extracted:
    """Plain text has no pages or headings, so a citation names line numbers."""
    lines = _read_text_file(path).splitlines()
    out = Extracted("txt", title=path.stem)
    paras: list[tuple[int, int, str]] = []          # (first line, last line, text)
    start, cur = None, []
    for i, line in enumerate(lines + [""], start=1):
        if line.strip():
            if start is None:
                start = i
            cur.append(line.strip())
        elif cur:
            paras.append((start, i - 1, " ".join(cur)))
            start, cur = None, []
    group: list[tuple[int, int, str]] = []
    n = 0

    def flush():
        if group:
            a, b = group[0][0], group[-1][1]
            loc = f"line {a}" if a == b else f"lines {a}–{b}"
            text = " ".join(p[2] if p[2].endswith((".", "!", "?", ":")) else p[2] + "." for p in group)
            out.sections.append(Section(0, loc, "", _clean_ws(text)))
        group.clear()

    for p in paras:
        w = _word_count(p[2])
        if group and n + w > TXT_SECTION_WORDS:
            flush()
            n = 0
        group.append(p)
        n += w
    flush()
    return out


EXTRACTORS = {"pdf": _extract_pdf, "pptx": _extract_pptx, "docx": _extract_docx,
              "md": _extract_md, "txt": _extract_txt}


def extract(path) -> Extracted:
    p = Path(path)
    kind = SUPPORTED.get(p.suffix.lower())
    if kind is None:
        raise ValueError(f"unsupported file type: {p.suffix or '(none)'}")
    return EXTRACTORS[kind](p)


# ------------------------------------------------------------------ embedders
class Embedder(Protocol):
    name: str
    dim: int

    def encode(self, texts: list[str], is_query: bool = False) -> np.ndarray: ...


BGE_NAME = "BAAI/bge-small-en-v1.5"


def bge_cached() -> bool:
    """Is bge-small in the local Hugging Face cache? (Never downloads.)"""
    try:
        from huggingface_hub import try_to_load_from_cache
        hit = try_to_load_from_cache(BGE_NAME, "config.json")
        return isinstance(hit, str) and os.path.exists(hit)
    except Exception:
        return False


class BgeEmbedder:
    """bge-small-en-v1.5 on the CPU, exactly as yaadash_v2 uses it (same model, same
    query instruction, normalised so cosine == dot). Loaded with
    local_files_only=True — the HF_HUB_OFFLINE behaviour — so a missing model is an
    error here, never a download. When yaadash_v2 already holds the model in this
    process, that copy is reused instead of loading a second 130 MB one."""
    name = BGE_NAME
    dim = 384
    QUERY_PREFIX = "Represent this sentence for searching relevant passages: "

    def __init__(self):
        self._model = None
        self._lock = threading.Lock()

    def _get(self):
        if self._model is None:
            with self._lock:
                if self._model is None:
                    y = sys.modules.get("yaadash_v2")
                    shared = getattr(y, "_model", None) if y is not None else None
                    if shared is not None:
                        self._model = shared
                    else:
                        from sentence_transformers import SentenceTransformer
                        self._model = SentenceTransformer(self.name, device="cpu",
                                                          local_files_only=True)
        return self._model

    def encode(self, texts: list[str], is_query: bool = False) -> np.ndarray:
        if is_query:
            texts = [self.QUERY_PREFIX + t for t in texts]
        v = self._get().encode(list(texts), normalize_embeddings=True,
                               show_progress_bar=False, batch_size=32)
        return np.asarray(v, dtype=np.float32)


def default_embedder() -> Optional[Embedder]:
    return BgeEmbedder() if bge_cached() else None


# ------------------------------------------------------------------ the LLM
LLM = Callable[[str, str], str]          # (system, user) -> reply text


def ollama_llm(system: str, user: str) -> str:
    """The local model through Ollama. Raises RuntimeError when it is not running."""
    import httpx
    model = os.environ.get("JARVIS_LIBRARY_MODEL", DEFAULT_MODEL)
    try:
        r = httpx.post(f"{OLLAMA}/api/chat", timeout=120, json={
            "model": model, "stream": False, "keep_alive": "30m",
            "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}],
            # Low temperature: this is reading comprehension, not writing. num_ctx
            # 4096 holds five 250-word passages (~1.7k tokens) plus the rules.
            "options": {"temperature": 0.1, "num_ctx": 4096, "num_predict": 400}})
        r.raise_for_status()
        return r.json()["message"]["content"]
    except Exception as e:
        raise RuntimeError(f"local model unreachable ({type(e).__name__})") from e


def llm_name(fn: LLM) -> str:
    """What to show as the answer's author: the Ollama model, or what a caller
    injected (its `model_name`, else "injected")."""
    if fn is ollama_llm:
        return os.environ.get("JARVIS_LIBRARY_MODEL", DEFAULT_MODEL)
    return getattr(fn, "model_name", None) or "injected"


# Live run, 2026-10-12: asked "setup time aur hold time mein kya farak hai?", the
# model gave both definitions with correct citations and then appended the rule-2
# sentence word for word, contradicting itself. A rule forbidding that did not
# help (cited-and-refused in 4/50 answers with this prompt, 5/50 with the rule,
# 25 questions x 2), so this prompt stays and parse_answer() deals with it.
SYSTEM_PROMPT = f"""You answer questions using ONLY the numbered passages, which come from the user's own study notes.
Rules:
1. After every sentence that states a fact, put the number of the passage it came from in square brackets, like [2]. If two passages support it, write [1][3].
2. Use nothing you know from anywhere else, even if you are sure it is true. If the passages do not contain the answer, reply with exactly this sentence and nothing else: {NOT_IN_NOTES}
3. If the passages answer only part of the question, answer that part with citations, then say which part the notes do not cover.
4. Copy numbers, part numbers, pin numbers and formulas exactly as the passage gives them.
5. At most five sentences. No preamble."""


def build_prompt(question: str, sources: list[dict]) -> tuple[str, str]:
    """(system, user) for the LLM. Passage n is sources[n-1]; its header names the
    file and location so the model can also say it in words."""
    blocks = []
    for s in sources:
        blocks.append(f"[{s['n']}] {s['file']}, {s['loc']}\n{s['text']}")
    user = "PASSAGES\n\n" + "\n\n".join(blocks) + f"\n\nQUESTION: {question.strip()}"
    return SYSTEM_PROMPT, user


_CITE = re.compile(r"\[(\d+(?:\s*(?:,|–|-|and)\s*\d+)*)\]")
_DONT_KNOW = re.compile(r"\b(i\s+don'?t\s+know|i\s+do\s+not\s+know|(?:notes|passages)\s+(?:don'?t|do\s+not)\s+cover)\b", re.I)


def cited_numbers(text: str) -> list[int]:
    """Every passage number cited in `text`, in order of first use: [1], [1, 3],
    [2-4] and [1][2] all count."""
    out: list[int] = []
    for m in _CITE.finditer(text or ""):
        body = m.group(1)
        for a, b in re.findall(r"(\d+)\s*(?:–|-)\s*(\d+)", body):
            for n in range(int(a), int(b) + 1):
                if n not in out:
                    out.append(n)
        body = re.sub(r"(\d+)\s*(?:–|-)\s*(\d+)", "", body)
        for n in re.findall(r"\d+", body):
            if int(n) not in out:
                out.append(int(n))
    return out


_STRAY_REFUSAL = re.compile(r"\s*I\s+don'?t\s+know\s*[—–-]+\s*your\s+notes\s+don'?t\s+cover\s+that\.?", re.I)
CAVEAT_NOT_FULLY = "The model also said your notes do not fully cover this question."


def parse_answer(text: str, n_sources: int) -> dict:
    """What the model's reply cites, and whether it abstained or failed to cite.

    A cited answer that ALSO carries the refusal sentence contradicts itself. The
    sentence is moved out of the answer into `caveat`, which the UI shows, rather
    than deleted: the model may have meant that part of the question is missing."""
    text = (text or "").strip()
    nums = cited_numbers(text)
    used = [n for n in nums if 1 <= n <= n_sources]
    invalid = [n for n in nums if not 1 <= n <= n_sources]
    caveat = ""
    if used and _STRAY_REFUSAL.search(text):
        text = _STRAY_REFUSAL.sub("", text).strip()
        caveat = CAVEAT_NOT_FULLY
    abstained = not used and bool(_DONT_KNOW.search(text))
    return {"answer": text, "used": used, "invalid": invalid, "abstained": abstained,
            "uncited": not used and not abstained and bool(text), "caveat": caveat}


def speakable(result: dict) -> str:
    """The answer as JARVIS should SAY it: no "[1][2]" read aloud, the sources named
    once at the end ("From your notes: DE_L4_Sequential, page 3.")."""
    ans = result.get("answer") or ""
    if result.get("abstained") or not ans:
        return ans or NOT_IN_NOTES
    spoken = _WS.sub(" ", _CITE.sub("", ans)).replace(" .", ".").strip()
    names = []
    by_n = {s["n"]: s for s in result.get("sources", [])}
    for n in result.get("used", []):
        s = by_n.get(n)
        if not s:
            continue
        loc = s["loc"].replace("p. ", "page ").replace("§ ", "").replace(" › ", ", ").replace(" · ", ", ")
        name = f"{Path(s['file']).stem.replace('_', ' ')}, {loc}"
        if name not in names:
            names.append(name)
    if names:
        spoken += " From your notes: " + "; ".join(names[:3]) + "."
    return spoken


# ------------------------------------------------------------------ the library
@dataclass
class Retrieval:
    hits: list[dict]
    abstained: bool
    reason: str
    best_distance: Optional[float]
    coverage: float
    mode: str
    took_ms: int


# Abstention gate, measured 2026-10-12 on the synthetic semester in
# yaadash/tests/test_library_eval.py (38 passages; bge-small cosine distance of the
# best passage; coverage = IDF-weighted share of the question's words found in the
# top-5 passages):
#   answerable (20):          distance 0.186-0.346, coverage 0.56-1.00
#   off-topic (3):            distance 0.458-0.585, coverage 0.22-0.33
#   right topic, fact absent: distance 0.341 / 0.353 ("74HC00 propagation delay",
#                             "Verilog blocking"), coverage 0.38 / 0.22
# Distance alone cannot separate the last group from real questions (0.341 < 0.346),
# so the gate has two rules: far away -> refuse (0.40, midway between 0.346 and
# 0.458); or most of the question's words missing AND only a middling match ->
# refuse. What slips through is left to the prompt's "I don't know" rule, which the
# live model followed 5/5 when the gate was bypassed (test_live_model_citations).
# Hinglish: bge-small is English-only, so "K-map mein ones ko group kaise karte
# hain?" sat at distance 0.42 and was refused, although k, map and group are on
# the right page (coverage 0.69; the page says "1s", not "ones"). Word evidence
# over at least two words with coverage >= 0.60 now overrides the distance rule up
# to 0.50. The best-covered unanswerable question has 0.38. A single shared word is
# not evidence: "what time is it?" (coverage 1.00 on "time", distance 0.42) is still
# refused. I set this after seeing the miss, and all of these numbers come from my
# invented notes: re-measure on real ones before trusting them.
GATE_MAX_DISTANCE = 0.40
GATE_HARD_DISTANCE = 0.50     # beyond this nothing overrides the refusal
GATE_MIN_COVERAGE = 0.45
GATE_SOFT_DISTANCE = 0.30
GATE_STRONG_COVERAGE = 0.60


def _bm25(corpus: list[list[str]]):
    """BM25Okapi (as in yaadash_recall) with Lucene's IDF, log(1 + (N - n + .5)/(n + .5)).

    Okapi's own IDF is zero for a word in exactly half the passages and negative
    above that. yaadash_recall never sees it (110k messages), but a library starts
    small: with two documents ingested, "compressor warranty AC" scored 0 against
    every passage and nothing was found. Lucene's form also falls as a word gets
    commoner, but it flattens towards zero instead of crossing it."""
    import math
    from rank_bm25 import BM25Okapi

    class _LuceneIdf(BM25Okapi):
        def _calc_idf(self, nd):
            for word, freq in nd.items():
                self.idf[word] = math.log(1 + (self.corpus_size - freq + 0.5) / (freq + 0.5))
            self.average_idf = sum(self.idf.values()) / max(len(self.idf), 1)

    return _LuceneIdf(corpus)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _norm(path) -> str:
    return os.path.normcase(os.path.abspath(str(path)))


class Library:
    """One study library on disk. Thread-safe for one writer and many readers."""

    def __init__(self, root=None, embedder="auto", llm: Optional[LLM] = None):
        self.root = Path(root) if root else default_root()
        self.root.mkdir(parents=True, exist_ok=True)
        self.embedder: Optional[Embedder] = default_embedder() if embedder == "auto" else embedder
        self.embedder_error: Optional[str] = None
        self.llm: LLM = llm or ollama_llm
        self._lock = threading.RLock()
        self._db = None
        self._manifest: Optional[dict] = None
        self._manifest_mtime: Optional[float] = None
        self._corpus = None                         # (generation, ids, rows, bm25)

    # ------------------------------------------------------------ manifest
    @property
    def manifest_path(self) -> Path:
        return self.root / MANIFEST

    def _load(self) -> dict:
        """The manifest, re-read when another process changed it."""
        p = self.manifest_path
        try:
            mt = p.stat().st_mtime
        except FileNotFoundError:
            if self._manifest is None:
                self._manifest = {"version": 1, "generation": 0, "docs": {}}
            return self._manifest
        if self._manifest is None or mt != self._manifest_mtime:
            self._manifest = json.loads(p.read_text(encoding="utf-8"))
            self._manifest_mtime = mt
        return self._manifest

    def _save(self, m: dict) -> None:
        m["generation"] = int(m.get("generation", 0)) + 1
        tmp = self.manifest_path.with_suffix(".tmp")
        tmp.write_text(json.dumps(m, indent=1, ensure_ascii=False), encoding="utf-8")
        os.replace(tmp, self.manifest_path)          # atomic: a reader never sees half
        self._manifest = m
        self._manifest_mtime = self.manifest_path.stat().st_mtime

    # ------------------------------------------------------------ table
    def _connect(self):
        if self._db is None:
            import lancedb
            # Strong consistency: an ask right after an ingest must see the new rows.
            self._db = lancedb.connect(str(self.root / "lancedb"),
                                       read_consistency_interval=timedelta(0))
        return self._db

    def _table(self, create_dim: Optional[int] = None):
        db = self._connect()
        try:
            return db.open_table(TABLE)
        except Exception:
            if create_dim is None:
                return None
        import pyarrow as pa
        schema = pa.schema([
            pa.field("id", pa.string()), pa.field("doc_id", pa.string()),
            pa.field("file", pa.string()), pa.field("kind", pa.string()),
            pa.field("page", pa.int32()), pa.field("loc", pa.string()),
            pa.field("heading", pa.string()), pa.field("chunk", pa.int32()),
            pa.field("text", pa.string()), pa.field("index_text", pa.string()),
            pa.field("vector", pa.list_(pa.float32(), create_dim)),
        ])
        return db.create_table(TABLE, schema=schema)

    @property
    def dim(self) -> int:
        return self.embedder.dim if self.embedder is not None else 384

    def _embed(self, texts: list[str], is_query: bool = False) -> Optional[np.ndarray]:
        """Vectors, or None when no embedder works (the library then runs on BM25
        alone and says so in every answer)."""
        if self.embedder is None or not texts:
            return None
        try:
            return self.embedder.encode(texts, is_query=is_query)
        except Exception as e:
            self.embedder_error = f"{type(e).__name__}: {e}"[:300]
            self.embedder = None
            return None

    # ------------------------------------------------------------ ingest
    def _doc_by_path(self, m: dict, path: Path) -> Optional[dict]:
        key = _norm(path)
        for d in m["docs"].values():
            if _norm(d["path"]) == key or key in (_norm(p) for p in d.get("also_at", [])):
                return d
        return None

    def ingest_file(self, path) -> dict:
        """Add or refresh one file. Unchanged files are a no-op: size and mtime are
        checked first, then the content hash, so touching a file does not re-embed it."""
        p = Path(path).resolve()
        res = {"file": p.name, "path": str(p)}
        kind = SUPPORTED.get(p.suffix.lower())
        if kind is None:
            return {**res, "status": "skipped", "note": "unsupported type"}
        try:
            st = p.stat()
        except OSError as e:
            return {**res, "status": "failed", "note": str(e)}
        if st.st_size > MAX_FILE_BYTES:
            return {**res, "status": "skipped", "note": "larger than 200 MB"}
        with self._lock:
            m = self._load()
            old = self._doc_by_path(m, p)
            if (old and old["size"] == st.st_size and old["mtime"] == st.st_mtime
                    and _norm(old["path"]) == _norm(p)):
                return {**res, "status": "unchanged", "doc_id": old["doc_id"]}
        sha = _sha256(p)
        doc_id = sha[:16]
        with self._lock:
            m = self._load()
            old = self._doc_by_path(m, p)
            if old and old["sha256"] == sha:
                if _norm(old["path"]) == _norm(p):
                    old["size"], old["mtime"] = st.st_size, st.st_mtime
                    self._save(m)
                return {**res, "status": "unchanged", "doc_id": old["doc_id"]}
            same = m["docs"].get(doc_id)
            if same is not None and _norm(same["path"]) != _norm(p):
                if Path(same["path"]).exists():
                    also = same.setdefault("also_at", [])
                    if str(p) not in also:
                        also.append(str(p))
                        self._save(m)
                    return {**res, "status": "duplicate", "doc_id": doc_id,
                            "note": f"same content as {same['file']}"}
                # The file moved: keep its passages, follow it to the new path.
                same.update(path=str(p), file=p.name, size=st.st_size, mtime=st.st_mtime,
                            updated=_now())
                self._save(m)
                tbl = self._table()
                if tbl is not None:
                    tbl.update(where=f"doc_id = '{doc_id}'", values={"file": p.name})
                return {**res, "status": "moved", "doc_id": doc_id}

        try:
            ex = extract(p)
        except Exception as e:
            return {**res, "status": "failed", "note": f"could not read: {type(e).__name__}: {e}"[:300]}
        rows = []
        for sec in ex.sections:
            for piece in chunk_text(sec.text):
                head = f"{ex.title} — {sec.heading}" if sec.heading else ex.title
                rows.append({"id": f"{doc_id}:{len(rows):04d}", "doc_id": doc_id, "file": p.name,
                             "kind": kind, "page": int(sec.page), "loc": sec.loc,
                             "heading": sec.heading, "chunk": len(rows), "text": piece,
                             # The title and heading ride along for search only: a
                             # passage that never repeats "Karnaugh" still belongs to it.
                             "index_text": f"{head}\n{piece}"})
        vecs = self._embed([r["index_text"] for r in rows]) if rows else None
        emb_name = self.embedder.name if (self.embedder is not None and vecs is not None) else None
        for i, r in enumerate(rows):
            r["vector"] = (vecs[i] if vecs is not None else np.zeros(self.dim, np.float32)).tolist()

        with self._lock:
            m = self._load()
            old = self._doc_by_path(m, p)
            dim = len(rows[0]["vector"]) if rows else self.dim
            tbl = self._table(create_dim=dim)
            if self._table_dim(tbl) != dim:
                # The table was built by an embedder with another dimension.
                self.sync_vectors()
                tbl = self._table(create_dim=dim)
            if old is not None:
                # The old content may still live in a duplicate elsewhere: then that
                # copy keeps the passages and only this path moves on.
                others = [a for a in old.get("also_at", []) if _norm(a) != _norm(p) and Path(a).exists()]
                if _norm(old["path"]) != _norm(p):
                    old["also_at"] = [a for a in old.get("also_at", []) if _norm(a) != _norm(p)]
                elif others:
                    old.update(path=others[0], file=Path(others[0]).name,
                               also_at=[a for a in others[1:]])
                    tbl.update(where=f"doc_id = '{old['doc_id']}'", values={"file": old["file"]})
                else:
                    tbl.delete(f"doc_id = '{old['doc_id']}'")
                    m["docs"].pop(old["doc_id"], None)
            if rows:
                tbl.add(rows)
            m["docs"][doc_id] = {
                "doc_id": doc_id, "path": str(p), "file": p.name, "kind": kind,
                "title": ex.title, "sha256": sha, "size": st.st_size, "mtime": st.st_mtime,
                "pages": ex.pages, "chunks": len(rows),
                "words": sum(_word_count(s.text) for s in ex.sections),
                "scanned": ex.scanned, "blank": ex.blank, "ocr": ex.ocr, "notes": ex.notes,
                "embedder": emb_name, "added": (old or {}).get("added") or _now(), "updated": _now(),
            }
            self._save(m)
        status = "updated" if old is not None else "added"
        return {**res, "status": status, "doc_id": doc_id, "pages": ex.pages,
                "chunks": len(rows), "scanned": len(ex.scanned), "note": "; ".join(ex.notes)}

    def files_in(self, folder) -> list[Path]:
        """Supported files under `folder`, without following links (a junction loop
        must not hang the ingest) and without app/tool folders."""
        out: list[Path] = []
        for dirpath, dirnames, filenames in os.walk(folder, followlinks=False):
            dirnames[:] = sorted(d for d in dirnames
                                 if d.lower() not in SKIP_DIRS and not d.startswith("."))
            for f in sorted(filenames):
                if f.startswith(("~$", ".")):          # Office lock files, dotfiles
                    continue
                if Path(f).suffix.lower() in SUPPORTED:
                    out.append(Path(dirpath) / f)
                    if len(out) >= MAX_FILES:
                        return out
        return out

    def ingest(self, path, progress: Optional[Callable[[dict], None]] = None) -> dict:
        """Ingest a file or a folder. `progress` gets {done, total, current, last}."""
        p = Path(path).expanduser()
        if p.is_file():
            files = [p]
        elif p.is_dir():
            files = self.files_in(p)
        else:
            raise FileNotFoundError(f"no such file or folder: {p}")
        t0 = time.time()
        results = []
        for i, f in enumerate(files):
            if progress:
                progress({"done": i, "total": len(files), "current": f.name,
                          "last": results[-1] if results else None})
            results.append(self.ingest_file(f))
        if progress:
            progress({"done": len(files), "total": len(files), "current": None,
                      "last": results[-1] if results else None})
        counts: dict[str, int] = {}
        for r in results:
            counts[r["status"]] = counts.get(r["status"], 0) + 1
        self.sync_vectors()
        return {"path": str(p), "files": len(files), "counts": counts, "results": results,
                "took_s": round(time.time() - t0, 2)}

    @staticmethod
    def _table_dim(tbl) -> int:
        return tbl.schema.field("vector").type.list_size

    def sync_vectors(self) -> int:
        """Embed passages stored without vectors (ingested while no embedder worked)
        or by another embedder. Returns the number of documents re-embedded."""
        if self.embedder is None:
            return 0
        import pyarrow as pa
        import pyarrow.compute as pc
        with self._lock:
            m = self._load()
            tbl = self._table()
            name = self.embedder.name
            stale = [d for d in m["docs"].values() if d.get("embedder") != name and d["chunks"]]
            if tbl is None or not stale:
                return 0
            rebuild = self._table_dim(tbl) != self.embedder.dim
            if rebuild:                     # a new dimension needs a new table
                stale = [d for d in m["docs"].values() if d["chunks"]]
            ids = [d["doc_id"] for d in stale]
            quoted = ",".join(f"'{i}'" for i in ids)
            every = tbl.to_arrow()
            rows = every.filter(pc.is_in(every["doc_id"], value_set=pa.array(ids))).to_pylist()
            vecs = self._embed([r["index_text"] for r in rows])
            if vecs is None:
                return 0
            for r, v in zip(rows, vecs):
                r["vector"] = v.tolist()
            if rebuild:
                self._connect().drop_table(TABLE)
                tbl = self._table(create_dim=vecs.shape[1])
            else:
                tbl.delete(f"doc_id IN ({quoted})")
            if rows:
                tbl.add(rows)
            for d in stale:
                d["embedder"] = name
            self._save(m)
            return len(stale)

    def remove(self, doc_id: str) -> bool:
        if not re.fullmatch(r"[0-9a-f]{16}", doc_id or ""):
            return False
        with self._lock:
            m = self._load()
            if doc_id not in m["docs"]:
                return False
            tbl = self._table()
            if tbl is not None:
                tbl.delete(f"doc_id = '{doc_id}'")
            m["docs"].pop(doc_id)
            self._save(m)
            return True

    # ------------------------------------------------------------ views
    def documents(self) -> list[dict]:
        m = self._load()
        keep = ("doc_id", "file", "path", "kind", "title", "pages", "chunks", "words",
                "scanned", "blank", "ocr", "notes", "embedder", "added", "updated", "also_at")
        docs = [{k: d.get(k) for k in keep} for d in m["docs"].values()]
        return sorted(docs, key=lambda d: d["file"].lower())

    def status(self) -> dict:
        docs = self.documents()
        return {
            "root": str(self.root),
            "documents": len(docs),
            "chunks": sum(d["chunks"] or 0 for d in docs),
            "pages": sum(d["pages"] or 0 for d in docs),
            "scanned_pages": sum(len(d["scanned"] or []) for d in docs),
            "embedder": {"name": self.embedder.name if self.embedder else None,
                         "ready": self.embedder is not None, "error": self.embedder_error},
            "mode": "hybrid" if self.embedder is not None else "keyword-only",
            "ocr": ocr_available(),
        }

    # ------------------------------------------------------------ search
    def _load_corpus(self):
        """All passages (no vectors) + a BM25 index, rebuilt when the library changes."""
        m = self._load()
        gen = m.get("generation", 0)
        if self._corpus is not None and self._corpus[0] == gen:
            return self._corpus
        tbl = self._table()
        rows = []
        if tbl is not None and tbl.count_rows():
            cols = ["id", "doc_id", "file", "kind", "page", "loc", "heading", "chunk", "text", "index_text"]
            rows = tbl.to_arrow().select(cols).to_pylist()
        ids = [r["id"] for r in rows]
        bm = _bm25([tokens(r["index_text"]) for r in rows]) if rows else None
        self._corpus = (gen, ids, {r["id"]: r for r in rows}, bm)
        return self._corpus

    def _bm25_ranking(self, query: str, n: int = CANDIDATES) -> list[str]:
        _, ids, _, bm = self._load_corpus()
        if bm is None:
            return []
        q = query_terms(query)
        scores = bm.get_scores(q)
        order = np.argsort(scores)[::-1][:n]
        return [ids[i] for i in order if scores[i] > 0]

    def _vector_ranking(self, query: str, n: int = CANDIDATES) -> tuple[list[str], dict]:
        tbl = self._table()
        if tbl is None or self.embedder is None:
            return [], {}
        qv = self._embed([query], is_query=True)
        if qv is None:
            return [], {}
        m = self._load()
        ok = {d["doc_id"] for d in m["docs"].values() if d.get("embedder") == self.embedder.name}
        if not ok:
            return [], {}
        # _distance is named in the projection: leaving it implicit makes Lance log a
        # deprecation warning on every query.
        res = (tbl.search(qv[0].tolist()).distance_type("cosine")
               .select(["id", "doc_id", "_distance"]).limit(n * 2).to_list())
        res = [r for r in res if r["doc_id"] in ok][:n]
        return [r["id"] for r in res], {r["id"]: float(r["_distance"]) for r in res}

    def _coverage(self, query: str, hit_ids: list[str]) -> float:
        """IDF-weighted share of the question's words that appear in these passages.
        A word the library has never seen counts at the highest weight: "ethanol"
        missing from electronics notes is the strongest sign they cannot answer."""
        _, _, rows, bm = self._load_corpus()
        q = list(dict.fromkeys(query_terms(query)))
        if not q or bm is None:
            return 0.0
        idf_max = max(bm.idf.values()) if bm.idf else 1.0
        seen: set[str] = set()
        for i in hit_ids:
            seen.update(tokens(rows[i]["index_text"]))
        w = {t: max(bm.idf.get(t, idf_max), 0.05) for t in q}
        return sum(w[t] for t in q if t in seen) / sum(w.values())

    def search(self, query: str, k: int = 5, mode: str = "hybrid") -> list[dict]:
        """Ranked passages, no gate. mode: hybrid | bm25 | vector (for evaluation)."""
        return self._rank(query, k, mode)[0]

    def _rank(self, query: str, k: int, mode: str):
        _, _, rows, _ = self._load_corpus()
        bm = self._bm25_ranking(query) if mode in ("hybrid", "bm25") else []
        vec, dist = self._vector_ranking(query) if mode in ("hybrid", "vector") else ([], {})
        scores: dict[str, float] = {}
        for ranking in (bm, vec):
            for r, cid in enumerate(ranking):
                scores[cid] = scores.get(cid, 0.0) + 1.0 / (RRF_K + r + 1)
        order = sorted(scores, key=lambda c: -scores[c])
        top = order[:k]
        # Lexical guarantee (as in yaadash_recall): BM25's best passage always makes
        # the cut. A question naming a part number ("74LS47") is answered by the one
        # passage that has it, which vectors rank as merely "about displays".
        if mode == "hybrid" and bm and bm[0] not in top and k > 0:
            top = top[:k - 1] + [bm[0]]
        out = []
        for cid in top:
            r = rows.get(cid)
            if r is None:
                continue
            out.append({**{k2: r[k2] for k2 in ("id", "doc_id", "file", "kind", "page", "loc",
                                                 "heading", "chunk", "text")},
                        "score": round(scores.get(cid, 0.0), 5),
                        "bm25_rank": bm.index(cid) + 1 if cid in bm else None,
                        "vector_rank": vec.index(cid) + 1 if cid in vec else None,
                        "distance": round(dist[cid], 4) if cid in dist else None})
        best = min(dist.values()) if dist else None
        return out, best

    def retrieve(self, query: str, k: int = 5) -> Retrieval:
        t0 = time.time()
        mode = "hybrid" if self.embedder is not None else "keyword-only"
        q = (query or "").strip()
        if not q:
            return Retrieval([], True, "empty question", None, 0.0, mode, 0)
        _, ids, _, _ = self._load_corpus()
        if not ids:
            return Retrieval([], True, "the library is empty", None, 0.0, mode,
                             round((time.time() - t0) * 1000))
        hits, best = self._rank(q, k, "hybrid")
        cov = self._coverage(q, [h["id"] for h in hits])
        # The distance thresholds were measured on bge-small; another embedder's
        # distances live on another scale, so with it (or with none) only the word
        # coverage rule decides.
        d = best if (self.embedder is not None and self.embedder.name == BGE_NAME) else None
        strong_words = (cov >= GATE_STRONG_COVERAGE and len(set(query_terms(q))) >= 2
                        and (d is None or d <= GATE_HARD_DISTANCE))
        reason = ""
        if not hits:
            reason = "no passage shares a word with the question"
        elif d is not None and d > GATE_MAX_DISTANCE and not strong_words:
            reason = f"closest passage is too far (distance {d:.2f} > {GATE_MAX_DISTANCE})"
        elif cov < GATE_MIN_COVERAGE and (d is None or d > GATE_SOFT_DISTANCE):
            reason = f"the notes miss most of the question's words (coverage {cov:.2f})"
        return Retrieval(hits, bool(reason), reason, best, round(cov, 3), mode,
                         round((time.time() - t0) * 1000))

    def _sources(self, hits: list[dict]) -> list[dict]:
        m = self._load()
        out = []
        for i, h in enumerate(hits, start=1):
            d = m["docs"].get(h["doc_id"], {})
            out.append({"n": i, **h, "path": d.get("path"), "title": d.get("title")})
        return out

    def ask(self, question: str, k: int = 5, llm: Optional[LLM] = None) -> dict:
        """Answer from the notes with [n] citations, or say the notes do not cover it.

        The model is never asked when the gate abstains: an answer with no passages
        behind it is exactly what this exists to prevent. When the model is down,
        the passages are still returned (with llm_error) so the question is not
        wasted."""
        t0 = time.time()
        r = self.retrieve(question, k)
        sources = self._sources(r.hits)
        base = {"question": (question or "").strip(), "mode": r.mode, "sources": sources,
                "best_distance": r.best_distance, "coverage": r.coverage,
                "retrieval_ms": r.took_ms}
        if r.abstained:
            return {**base, "answer": NOT_IN_NOTES, "abstained": True, "reason": r.reason,
                    "used": [], "invalid": [], "uncited": False, "caveat": "", "llm": None,
                    "took_ms": round((time.time() - t0) * 1000)}
        system, user = build_prompt(question, sources)
        try:
            text = (llm or self.llm)(system, user)
        except Exception as e:
            return {**base, "answer": None, "abstained": False, "reason": "",
                    "used": [], "invalid": [], "uncited": False, "caveat": "", "llm": None,
                    "llm_error": str(e)[:300], "took_ms": round((time.time() - t0) * 1000)}
        parsed = parse_answer(text, len(sources))
        return {**base, **parsed, "reason": "the model said the passages do not answer it"
                if parsed["abstained"] else "", "llm": llm_name(llm or self.llm),
                "took_ms": round((time.time() - t0) * 1000)}


def context_block(lib: Library, question: str, k: int = 4) -> str:
    """For the voice brain: the passages as a prompt block, or "" when the gate
    abstains (so an unrelated question never drags lecture notes into the reply)."""
    r = lib.retrieve(question, k)
    if r.abstained or not r.hits:
        return ""
    src = lib._sources(r.hits)
    return ("FROM HIS STUDY NOTES (cite as 'your notes, <file> <page>'; if these do not "
            "answer the question, say his notes do not cover it):\n"
            + "\n".join(f"- [{s['file']}, {s['loc']}] {s['text']}" for s in src))


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="Study library: ingest notes, ask with citations.")
    ap.add_argument("--root", help="library folder (default data/yaadash/library)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("ingest")
    a.add_argument("path")
    b = sub.add_parser("ask")
    b.add_argument("question", nargs="+")
    sub.add_parser("list")
    args = ap.parse_args()
    lib = Library(args.root)
    if args.cmd == "ingest":
        print(json.dumps(lib.ingest(args.path, progress=lambda p: print(
            f"  {p['done']}/{p['total']} {p['current'] or ''}")), indent=1)[:4000])
    elif args.cmd == "list":
        for d in lib.documents():
            print(f"{d['file']:40s} {d['pages'] or '-':>4} pages {d['chunks']:>4} passages"
                  + (f"  ({len(d['scanned'])} scanned)" if d["scanned"] else ""))
    else:
        res = lib.ask(" ".join(args.question))
        print(res.get("answer") or f"[model unavailable: {res.get('llm_error')}]")
        for s in res["sources"]:
            mark = "*" if s["n"] in res.get("used", []) else " "
            print(f" {mark}[{s['n']}] {s['file']}, {s['loc']}")
