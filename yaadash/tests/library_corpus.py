"""A synthetic semester for testing the study library.

Everything here is INVENTED for the tests: four lecture PDFs on digital electronics
(one page of them a scan with no text layer), a lab manual in Word, revision notes in
Markdown, a slide deck, and one piece of household paperwork (an AC warranty card) —
the same pipeline is meant to answer "AC ki warranty kab tak hai?". No real notes, no
real home, no real people.

QUESTIONS were written before the library was tuned, each with the page (or heading)
that answers it; UNANSWERABLE are questions the notes do not cover — three plainly
off-topic, two on the right subject but asking for facts the notes do not hold.
"""
from __future__ import annotations

from pathlib import Path

# ----------------------------------------------------------------- lecture PDFs
# PDF text uses the base-14 Helvetica font, so it stays ASCII.
L1 = ("DE Lecture 1 - Number Systems and Codes", [
    """Positional number systems. In a positional system the value of a digit depends on its position: each place carries a weight that is a power of the base, or radix. Decimal uses base 10, binary uses base 2, octal base 8 and hexadecimal base 16. To find the decimal value of a binary number, multiply each bit by its weight (1, 2, 4, 8, 16 and so on from the right) and add the results; 101101 is 32 + 8 + 4 + 1 = 45. To convert a decimal integer into binary, divide it by 2 again and again and write down the remainders; reading the remainders from the last one to the first gives the binary number. For a fraction, multiply by 2 repeatedly and collect the integer parts instead. Hexadecimal is a compact way to write binary: group the bits in fours from the binary point, so 1011 0110 becomes B6. Octal groups the bits in threes. Engineers use hexadecimal for memory addresses and register values because every hex digit maps to exactly four bits.""",
    """Signed numbers. Sign-magnitude form uses the most significant bit as the sign and the remaining bits as the size; it is easy to read but has two zeros, +0 and -0, which complicates hardware. One's complement forms a negative number by inverting every bit, and it still has two zeros. Two's complement is what real processors use: to negate a number, invert all the bits and then add one. An n-bit two's complement number covers the range from -2^(n-1) to 2^(n-1) - 1, so eight bits hold values from -128 to +127 and sixteen bits from -32768 to +32767. Addition and subtraction use the same adder circuit, which is the main reason two's complement won. Overflow happens when the result does not fit in n bits: it is detected when the carry into the sign bit differs from the carry out of the sign bit, or equivalently when two numbers of the same sign produce a result of the opposite sign. Sign extension copies the sign bit into the new upper bits when a number is widened.""",
    """Binary codes. Binary coded decimal (BCD) stores each decimal digit separately in four bits, so 59 is written 0101 1001. Only the patterns 0000 to 1001 are used; the six patterns 1010 to 1111 are invalid in BCD because no single decimal digit needs them. When two BCD digits are added and the four-bit sum is greater than 9, or a carry comes out of the group, the result must be corrected by adding 0110 (six) to skip over the six unused codes. The Gray code is a unit-distance code: successive values differ in exactly one bit, which prevents glitches when a value changes. It is used in rotary shaft encoders, in asynchronous FIFO pointers and in the row and column order of Karnaugh maps. To convert binary to Gray, keep the most significant bit and XOR each pair of neighbouring binary bits to get the remaining Gray bits. Excess-3 code adds three to each BCD digit and is self-complementing, which made it popular in early calculators.""",
    None,   # a scanned tutorial sheet: an image with no text layer
    # Distractor: same vocabulary as the questions above (binary, two's complement,
    # 8 bits), different facts.
    """Worked arithmetic examples. Binary addition works column by column like decimal addition: 0110 + 0111 = 1101, which is 6 + 7 = 13. Subtraction is done by adding the two's complement of the subtrahend. For 13 - 9 in 8 bits, write 13 as 00001101 and -9 as 11110111; their sum is 1 00000100, and discarding the carry out of the eighth bit leaves 00000100, which is 4. Hexadecimal addition carries at sixteen: 9 + 8 = 11 hex. Real numbers are stored in floating point: the IEEE 754 single-precision format has 1 sign bit, 8 exponent bits with a bias of 127, and 23 fraction bits, while double precision uses 11 exponent bits and 52 fraction bits. Floating-point addition first aligns the exponents, then adds the fractions, then normalises the result.""",
])

L2 = ("DE Lecture 2 - Boolean Algebra and Karnaugh Maps", [
    """Boolean algebra. A Boolean variable takes only the values 0 and 1, and the basic operations are AND, OR and NOT. The laws mirror ordinary algebra with a few surprises. Idempotent law: A + A = A and A.A = A. Absorption law: A + AB = A and A(A + B) = A. Distributive law also holds in the unusual form A + BC = (A + B)(A + C). De Morgan's theorems are the most useful tools for simplification: the complement of a product equals the sum of the complements, (AB)' = A' + B', and the complement of a sum equals the product of the complements, (A + B)' = A'B'. In words, to complement an expression you break the bar and change the operator. The duality principle says that every identity stays true if AND and OR are swapped and 0 and 1 are swapped. Truth tables give a complete but long description of a function; algebra gives a short one.""",
    """Canonical forms and Karnaugh maps. A minterm is a product term that contains every variable once, and a sum of minterms is the canonical sum-of-products form; maxterms and the product-of-sums form are the dual. A Karnaugh map draws the truth table as a grid in which neighbouring cells differ in only one variable, because the rows and columns follow Gray code order. A four-variable map has 16 cells. To simplify, circle the 1s in rectangular groups whose size is a power of two: 1, 2, 4, 8 or 16 cells. Make each group as large as possible, because a larger group eliminates more variables and gives a product term with fewer literals, and use as few groups as possible. Groups may overlap, and the map wraps around: cells on the left edge are adjacent to cells on the right edge, and the top row is adjacent to the bottom row. Each group becomes one product term, and the simplified expression is the OR of those terms.""",
    """Don't-care conditions and prime implicants. Some input combinations never occur, for example the six unused codes of a BCD input. Their outputs are marked X in the map and called don't-care conditions. A don't-care may be treated as a 1 when that makes a group larger and as a 0 otherwise; it never has to be covered. A prime implicant is a group that cannot be made any larger. An essential prime implicant is a prime implicant that covers at least one minterm that no other prime implicant covers, so it must appear in every minimal solution. The procedure is to first select all essential prime implicants and then cover the remaining 1s with as few extra prime implicants as possible. For functions of more than five or six variables, maps become impractical and the tabular Quine-McCluskey method is used instead, because it can be programmed.""",
    """Universal gates. NAND and NOR are called universal gates because each of them alone can implement any Boolean function. A NAND gate becomes an inverter when both of its inputs are tied together, since (AA)' = A'. An AND gate is a NAND followed by a NAND inverter, and an OR gate is built by inverting both inputs before a NAND, which is De Morgan's theorem in hardware. A two-level sum-of-products circuit converts directly to a NAND-NAND circuit with the same structure, and a product-of-sums circuit converts to NOR-NOR. An exclusive-OR gate needs four NAND gates. Because a single gate type can build everything, manufacturers can optimise one cell, and integrated circuits such as the quad two-input NAND chip are the workhorses of the lab.""",
    """Worked K-map example. Simplify F(A, B, C, D) = sum of minterms (0, 2, 5, 7, 8, 10, 13, 15). Plot the eight 1s on a four-variable map. The four corner cells 0, 2, 8 and 10 are adjacent because the map wraps around in both directions, so they form a group of four that gives the term B'D'. The four centre cells 5, 7, 13 and 15 form a second group of four that gives BD. Therefore F = BD + B'D', which is the exclusive-NOR of B and D. For a product-of-sums answer, group the 0s instead and write each group as a sum term. Always check the result against two or three rows of the truth table. A common mistake is to group cells diagonally, which is never allowed, because diagonal cells differ in two variables.""",
])

L3 = ("DE Lecture 3 - Combinational Circuits", [
    """Multiplexers. A multiplexer, or data selector, connects one of several data inputs to a single output. A multiplexer with 2^n data inputs needs n select lines, so a 4-to-1 multiplexer has two select lines and an 8-to-1 multiplexer has three. The select lines form a binary number that chooses which input is passed through. A multiplexer can implement any Boolean function: a function of n variables fits on a multiplexer with 2^(n-1) inputs if one variable, or its complement, or a constant 0 or 1, is applied to each data input while the other variables drive the select lines. A demultiplexer does the reverse, routing one input to one of many outputs. Multiplexers are used to share a bus among several sources and to build parallel-to-serial converters.""",
    """Decoders and encoders. A decoder converts an n-bit code into one of 2^n output lines, one line active at a time. The 74LS138 is a 3-to-8 line decoder with active-low outputs and three enable inputs, one active-high and two active-low, which makes it easy to cascade two of them into a 4-to-16 decoder; it is widely used for memory address decoding. An encoder does the opposite: it produces the binary code of the one active input. A plain encoder gives a wrong code if two inputs are active together, so a priority encoder is used instead. The 74LS148 is an 8-to-3 priority encoder: it outputs the code of the highest-numbered active input and has a separate output that says whether any input is active at all. BCD to seven-segment decoders drive numeric displays.""",
    """Adders and comparators. A half adder adds two bits: the sum is A XOR B and the carry is A AND B. A full adder adds two bits and an incoming carry, and can be built from two half adders and an OR gate. Connecting n full adders in a chain, with the carry out of each stage feeding the carry in of the next, gives a ripple carry adder. It is simple but slow: each stage must wait for the carry from the stage before it, so the worst-case delay grows linearly with the number of bits. A carry look-ahead adder removes the wait by computing, for each bit, a generate signal G = AB and a propagate signal P = A XOR B, and then producing all carries in parallel from these signals with two levels of logic. The price is more gates. A magnitude comparator decides whether one binary number is greater than, equal to or less than another.""",
    """Code converters and parity. A code converter changes one binary code into another with combinational logic. A binary-to-Gray converter needs only XOR gates, one for each bit after the most significant, and a Gray-to-binary converter chains the XOR gates from the top bit downwards. A BCD-to-excess-3 converter is designed from a truth table with don't-cares for the six unused inputs. A parity generator adds one bit so that the total number of 1s is even (even parity) or odd (odd parity); a tree of XOR gates computes it. At the receiver a parity checker repeats the calculation, and a mismatch shows that a single bit was corrupted, though parity cannot say which bit or detect two errors.""",
])

L4 = ("DE Lecture 4 - Sequential Circuits", [
    """Latches. A sequential circuit has memory: its outputs depend on past inputs as well as present ones. The simplest memory element is the SR latch, built from two cross-coupled NOR gates. Setting S to 1 makes Q equal to 1, setting R to 1 makes Q equal to 0, and with both inputs at 0 the latch holds its previous state. The input S = R = 1 is forbidden: it forces both Q and Q' to 0, which breaks the rule that they are complements, and when both inputs return to 0 together the final state depends on which gate is faster, a race condition. A gated or clocked latch adds an enable input, so the latch responds only while the enable is high; it is level-sensitive, which means that changes on the inputs pass straight through for the whole time the enable is active.""",
    """Flip-flops. A flip-flop is edge-triggered: it samples its inputs only at a rising or falling clock edge. The D flip-flop simply copies D to Q at the clock edge, and it is the basic storage element of registers. The JK flip-flop removes the forbidden state of the SR latch: J = 1, K = 0 sets, J = 0, K = 1 resets, J = K = 0 holds, and J = K = 1 makes the output toggle, changing to its complement at every clock edge. Its characteristic equation is Q+ = JQ' + K'Q. The T flip-flop toggles when T = 1 and holds when T = 0, which makes it the natural building block for counters. In a level-triggered JK latch with J = K = 1 the output can toggle many times while the clock is high, the race-around condition; the master-slave construction, or true edge triggering, removes it.""",
    """Timing of flip-flops. Setup time is the minimum time for which the data input must be held stable before the active clock edge. Hold time is the minimum time for which the data must stay stable after the clock edge. If either requirement is violated, the flip-flop may become metastable: its output hovers between logic levels for an unpredictable time before settling to 0 or 1. Signals that come from outside the clock domain, such as a push button, can change at any moment, so they must pass through a synchronizer, two flip-flops connected in series, which gives the first stage a full clock period to resolve. The clock-to-Q delay is the time from the edge to a valid output. The maximum clock frequency of a path is f_max = 1 / (t_cq + t_logic + t_setup), where t_logic is the delay of the combinational logic between the two flip-flops.""",
    """Counters. In a ripple, or asynchronous, counter each flip-flop is clocked by the output of the one before it, so the delays add up and the outputs settle one after another, producing brief false counts. In a synchronous counter every flip-flop shares the same clock and the next-state logic decides which ones toggle, so all outputs change together. A mod-N counter counts from 0 to N-1 and then returns to 0; it can be made by resetting a binary counter when it reaches N. The 74LS90 is a decade counter, mod-10, made of a divide-by-2 and a divide-by-5 section. A ring counter circulates a single 1 through a shift register and has as many states as flip-flops. A Johnson counter, or twisted ring counter, feeds back the complement of the last output and gives 2n states with n flip-flops.""",
    """Shift registers. A shift register is a chain of D flip-flops that moves its data one position at every clock pulse. They are classified by how data enters and leaves: serial in serial out (SISO), serial in parallel out (SIPO), parallel in serial out (PISO) and parallel in parallel out (PIPO). The 74LS194 is a 4-bit universal shift register that can shift left, shift right, load in parallel or hold, chosen by two mode inputs. Shift registers convert serial data to parallel and back, for example in a UART, delay a signal by a whole number of clock periods, and, with XOR feedback from selected taps, form a linear feedback shift register that generates a pseudo-random sequence for testing and scrambling.""",
    """Timing in practice. Clock skew is the difference between the times at which the same clock edge arrives at two flip-flops. Positive skew, where the capturing flip-flop sees the edge later, gives a path more time and helps with setup, but it eats into the hold margin. A setup violation can always be cured by slowing the clock down; a hold violation cannot, because it does not depend on the clock period, so it must be fixed by adding delay to the short path. Clock jitter is the cycle-to-cycle variation of the clock period and is subtracted from the timing budget. The 7474 dual D flip-flop chip has active-low preset and clear inputs that override the clock, which is useful to put a counter in a known state at power-up.""",
])

SCANNED_TEXT = "Tutorial sheet 1 - worked solutions (handwritten). Q1 convert 0.625 to binary."

# ----------------------------------------------------------------- lab manual (DOCX)
LAB_TITLE = "Digital Electronics Lab Manual"
LAB = [
    ("Experiment 1: Verification of logic gates", [
        "Aim: to verify the truth tables of the basic and universal gates using TTL integrated circuits.",
        "Components: 7400 quad two-input NAND, 7402 quad two-input NOR, 7404 hex inverter, 7408 AND, 7432 OR and 7486 XOR, a breadboard, a 5 V supply, LEDs and switches.",
        "All of these are 14-pin dual in-line packages with the same power pins: pin 14 is Vcc (+5 V) and pin 7 is ground. Connect power and ground before any input, and check the pin-out against the notch on the chip.",
        "Procedure: apply every input combination with the switches and record the LED state for each output in the observation table.",
    ]),
    ("Experiment 2: BCD to seven-segment decoder", [
        "Aim: to display a BCD digit on a common-anode seven-segment display using the 7447 decoder-driver.",
        "The 7447 has active-low open-collector outputs, so a segment lights when its output goes low. Put a 330 ohm current-limiting resistor in series with every segment; without it the LED and the chip are damaged.",
        "The lamp test input, when pulled low, lights all seven segments to check the display.",
    ]),
    None,   # a hard page break
    ("Experiment 3: Mod-10 counter using the 7490", [
        "Aim: to build a decade counter. Connect output QA to the clock input B so the divide-by-2 and divide-by-5 sections form one mod-10 counter.",
        "Ground both reset inputs R0(1) and R0(2) for normal counting. Clock the counter at 1 Hz from a 555 timer wired as an astable multivibrator and watch the count on LEDs.",
    ]),
    ("Experiment 4: Flip-flops with the 7474 and 7476", [
        "Aim: to verify the D and JK flip-flops. The 7474 holds two D flip-flops; the 7476 holds two JK flip-flops with preset and clear.",
        "The 7476 is a 16-pin chip with unusual power pins: pin 5 is Vcc and pin 13 is ground, so do not wire it like the 14-pin gate chips.",
        "Apply a single clock pulse from a debounced switch and record Q. With J and K both tied high the 7476 toggles on every falling edge of the clock.",
    ]),
    ("Safety and good practice", [
        "Switch off the supply before changing any wiring.",
        "Place a 0.1 uF decoupling capacitor between Vcc and ground close to each chip to absorb switching noise.",
        "Never leave unused TTL inputs floating: an open TTL input behaves like a high input but picks up noise, so tie unused inputs to Vcc through a 1 kilo-ohm resistor, or to ground when a low is required.",
        "Do not exceed 5.25 V on a TTL supply.",
    ]),
]
LAB_VIVA = [("Why is the NAND gate called universal?", "Any Boolean function can be built from NAND gates alone."),
            ("What is fan-out?", "The number of gate inputs one output can drive reliably.")]

# ----------------------------------------------------------------- revision notes (MD)
NOTES_MD = """---
title: Digital Electronics - exam revision
tags: [de, revision]
---
# FSM design
## Moore vs Mealy
- In a **Moore** machine the outputs depend only on the present state.
- In a **Mealy** machine the outputs depend on the present state *and* the inputs, so they react in the same clock cycle and usually need fewer states.

## State minimisation
Two states are equivalent when they give the same outputs and move to equivalent next states for every input; equivalent states are merged. The implication chart method finds them.

# Memories
## SRAM vs DRAM
An SRAM cell is a latch made of six transistors. It keeps its data for as long as power is applied, needs no refresh, and is fast, which is why caches use it.
A DRAM cell stores a bit as charge on a capacitor with one access transistor. The charge leaks away, so every row must be refreshed periodically, typically every 64 ms. DRAM is denser and cheaper per bit than SRAM, which is why main memory uses it.

## ROM types
PROM is programmed once. EPROM is erased with ultraviolet light through a quartz window. EEPROM is erased electrically one byte at a time, and Flash memory is erased electrically in whole blocks.

```verilog
// not part of the notes' prose: a code block must not become a heading
# not a heading
```

# Timing
## Static hazards
A static hazard is a brief glitch on an output that should stay constant, caused by unequal delays along two paths when one input changes. On the K-map it shows up as two adjacent groups that do not overlap; adding the redundant consensus term that covers both removes the hazard.

## Refresh and timing budgets
Every synchronous path must fit inside one clock period: clock-to-Q delay plus logic delay plus setup time. The memory controller also has to schedule refresh cycles between ordinary reads and writes.

# Logic families
## TTL vs CMOS
TTL (74LS) runs from 5 V plus or minus 5 percent and has a noise margin of about 0.4 V. CMOS draws almost no static power; its dynamic power grows with switching frequency as P = C V^2 f. See [[Lecture 4|sequential notes]] for timing.
"""

# ----------------------------------------------------------------- slides (PPTX)
SLIDES_TITLE = "DE Lecture 5 - Finite State Machines"
SLIDES = [
    ("Finite State Machines", ["Lecture 5", "Designing controllers with states"], ""),
    ("Example: traffic light controller", [
        "States: green, yellow, red for the main road",
        "A timer input moves the controller to the next state",
        "Outputs drive the three lamps"], "Draw the state diagram first, then the state table."),
    ("State encoding", [
        "Binary encoding uses the fewest flip-flops: ceil(log2 N) for N states",
        "One-hot encoding uses one flip-flop per state, with exactly one flip-flop set at a time",
        "One-hot needs more flip-flops but its next-state logic is simpler and faster, so FPGA tools prefer it"],
     "Gray encoding of states reduces switching when the sequence is fixed."),
    ("Recap: Moore and Mealy outputs", [
        "Moore: outputs are written inside the state bubbles",
        "Mealy: outputs are written on the transition arrows",
        "A Mealy output can glitch if its input glitches; register it if that matters"], ""),
]

# ----------------------------------------------------------------- household paperwork (TXT)
WARRANTY_TXT = """SPLIT AIR CONDITIONER 1.5 TON - WARRANTY CARD (sample, invented for tests)

Purchased on 14 March 2025, invoice number INV-2025-0412.

Comprehensive warranty on the whole unit: one year from the date of purchase, that is until 13 March 2026. It covers parts and labour.

Compressor warranty: ten years from the date of purchase, until 13 March 2035. It covers the compressor part only; gas charging and labour are charged.

PCB warranty: five years on the outdoor unit's control board.

Service: two free services within the first year. To claim warranty keep the invoice and this card. The serial number is printed on the right side of the indoor unit.
"""

# A second piece of paperwork that shares the warranty vocabulary (distractor).
INVERTER_TXT = """HOME INVERTER BATTERY 150 Ah - WARRANTY (sample, invented for tests)

Battery warranty: 36 months full replacement plus 12 months pro-rata from the date of installation, 2 June 2024.

Top up the battery with distilled water every two months and keep the terminals greased. The warranty is void if the battery is charged from an unapproved charger.
"""


# ----------------------------------------------------------------- the questions
# gold = (file, page) for PDFs and slides; (file, heading text) for DOCX/MD; (file, None) for TXT.
QUESTIONS = [
    ("How do I convert a decimal number into binary?", ("DE_L1_Number_Systems.pdf", 1)),
    ("What range of values fits in an 8-bit two's complement number?", ("DE_L1_Number_Systems.pdf", 2)),
    ("Why are the codes 1010 to 1111 not allowed in BCD, and what correction is applied after adding?", ("DE_L1_Number_Systems.pdf", 3)),
    ("Where is the Gray code used and how is it obtained from binary?", ("DE_L1_Number_Systems.pdf", 3)),
    ("State De Morgan's theorems.", ("DE_L2_Boolean_Kmaps.pdf", 1)),
    ("How should the ones be grouped when simplifying with a K-map?", ("DE_L2_Boolean_Kmaps.pdf", 2)),
    ("What makes a prime implicant essential?", ("DE_L2_Boolean_Kmaps.pdf", 3)),
    ("How can a NAND gate be turned into an inverter?", ("DE_L2_Boolean_Kmaps.pdf", 4)),
    ("How many select lines does a 4-to-1 mux need?", ("DE_L3_Combinational.pdf", 1)),
    ("What does the 74LS138 do?", ("DE_L3_Combinational.pdf", 2)),
    ("Why is a ripple carry adder slow, and how does carry look-ahead speed it up?", ("DE_L3_Combinational.pdf", 3)),
    ("What goes wrong in an SR latch when both inputs are 1?", ("DE_L4_Sequential.pdf", 1)),
    ("What does a JK flip-flop do when J and K are both high?", ("DE_L4_Sequential.pdf", 2)),
    ("What is the difference between setup time and hold time?", ("DE_L4_Sequential.pdf", 3)),
    ("How many states does a Johnson counter with n flip-flops go through?", ("DE_L4_Sequential.pdf", 4)),
    ("Which pin of the 7400 chip is ground?", ("Digital_Lab_Manual.docx", "Experiment 1")),
    ("What should I do with unused TTL inputs on the breadboard?", ("Digital_Lab_Manual.docx", "Safety and good practice")),
    ("Why does DRAM have to be refreshed when SRAM does not?", ("Exam_Revision_Notes.md", "SRAM vs DRAM")),
    ("What is one-hot state encoding and why do FPGA tools like it?", ("DE_L5_FSM_slides.pptx", 3)),
    ("When does the compressor warranty on the AC run out?", ("AC_warranty_card.txt", None)),
]

# The way I actually ask (Hinglish), scored separately from the 20 above. The
# embedder is English-only, so these lean on the English keywords inside them.
HINGLISH = [
    ("setup time aur hold time mein kya farak hai?", ("DE_L4_Sequential.pdf", 3)),
    ("8 bit two's complement mein kitni range hoti hai?", ("DE_L1_Number_Systems.pdf", 2)),
    ("7400 IC ka ground pin konsa hai?", ("Digital_Lab_Manual.docx", "Experiment 1")),
    ("AC ke compressor ki warranty kab tak hai?", ("AC_warranty_card.txt", None)),
    ("K-map mein ones ko group kaise karte hain?", ("DE_L2_Boolean_Kmaps.pdf", 2)),
]

UNANSWERABLE = [
    "What is the boiling point of ethanol?",
    "Who won the 2011 cricket world cup final?",
    "How do I apply for a German student visa?",
    "What is the propagation delay of a 74HC00 NAND gate in nanoseconds?",
    "Explain blocking and non-blocking assignments in Verilog.",
]


# ----------------------------------------------------------------- builders
def write_pdf(path: Path, title: str, pages: list, scanned_text: str = SCANNED_TEXT) -> Path:
    import fitz
    doc = fitz.open()
    for text in pages:
        page = doc.new_page(width=595, height=842)
        rect = fitz.Rect(56, 56, 539, 786)
        if text is None:
            # A "scan": the words exist only as pixels, as from a phone camera.
            tmp = fitz.open()
            tp = tmp.new_page(width=595, height=842)
            tp.insert_textbox(rect, scanned_text, fontsize=14, fontname="helv")
            pix = tp.get_pixmap(dpi=72)
            page.insert_image(page.rect, pixmap=pix)
            tmp.close()
            continue
        left = page.insert_textbox(rect, text, fontsize=10.5, fontname="helv")
        assert left >= 0, f"text overflowed a page of {path.name}"
    doc.set_metadata({"title": title})
    doc.save(str(path))
    doc.close()
    return path


def write_docx(path: Path) -> Path:
    import docx
    d = docx.Document()
    d.core_properties.title = LAB_TITLE
    d.add_heading(LAB_TITLE, level=0)
    for sec in LAB:
        if sec is None:
            d.add_page_break()
            continue
        head, paras = sec
        d.add_heading(head, level=1)
        for p in paras:
            d.add_paragraph(p)
    d.add_heading("Viva questions", level=1)
    t = d.add_table(rows=len(LAB_VIVA), cols=2)
    for row, (q, a) in zip(t.rows, LAB_VIVA):
        row.cells[0].text, row.cells[1].text = q, a
    d.save(str(path))
    return path


def write_pptx(path: Path) -> Path:
    from pptx import Presentation
    prs = Presentation()
    prs.core_properties.title = SLIDES_TITLE
    for i, (title, bullets, notes) in enumerate(SLIDES):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))
    return path


def build_corpus(folder: Path) -> Path:
    """Write the whole synthetic semester into `folder` (created) and return it."""
    folder = Path(folder)
    (folder / "lectures").mkdir(parents=True, exist_ok=True)
    (folder / "home").mkdir(exist_ok=True)
    write_pdf(folder / "lectures" / "DE_L1_Number_Systems.pdf", *L1)
    write_pdf(folder / "lectures" / "DE_L2_Boolean_Kmaps.pdf", *L2)
    write_pdf(folder / "lectures" / "DE_L3_Combinational.pdf", *L3)
    write_pdf(folder / "lectures" / "DE_L4_Sequential.pdf", *L4)
    write_pptx(folder / "lectures" / "DE_L5_FSM_slides.pptx")
    write_docx(folder / "Digital_Lab_Manual.docx")
    (folder / "Exam_Revision_Notes.md").write_text(NOTES_MD, encoding="utf-8")
    (folder / "home" / "AC_warranty_card.txt").write_text(WARRANTY_TXT, encoding="utf-8")
    (folder / "home" / "Inverter_battery_warranty.txt").write_text(INVERTER_TXT, encoding="utf-8")
    # Things an ingest must skip: an unsupported file, an Office lock file, a dot-folder.
    (folder / "circuit.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    (folder / "~$Digital_Lab_Manual.docx").write_bytes(b"lock")
    (folder / ".obsidian").mkdir(exist_ok=True)
    (folder / ".obsidian" / "workspace.md").write_text("# not notes\n", encoding="utf-8")
    return folder


def is_hit(hit: dict, gold: tuple) -> bool:
    """Does a retrieved passage come from the gold page / section?"""
    file, where = gold
    if hit["file"] != file:
        return False
    if where is None:
        return True
    if isinstance(where, int):
        return hit["page"] == where
    return where.lower() in (hit.get("heading") or hit.get("loc") or "").lower()
