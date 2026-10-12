"""The study library's parts, one behaviour at a time, on invented files.

Vectors here come from the hashing stand-in (conftest.HashingEmbedder) so the suite
stays fast; test_library_eval.py measures retrieval with the real bge-small.
"""
from __future__ import annotations

import os
import time
from pathlib import Path

import pytest

import library as L
import library_corpus as C


def words(s: str) -> int:
    return len(s.split())


def make_lib(tmp_path, embedder, llm=None, name="lib"):
    return L.Library(tmp_path / name, embedder=embedder, llm=llm)


# ------------------------------------------------------------------ chunking
def test_chunks_stay_near_the_target_and_overlap():
    sentences = [f"Sentence number {i} talks about flip-flops and their timing rules today." for i in range(120)]
    chunks = L.chunk_text(" ".join(sentences))
    assert len(chunks) >= 5
    assert all(words(c) <= L.MAX_WORDS + L.MIN_TAIL_WORDS for c in chunks)
    assert all(words(c) >= 150 for c in chunks[:-1])
    for a, b in zip(chunks, chunks[1:]):                 # the last sentence carries over
        assert a.split(". ")[-1].rstrip(".") in b
    # Nothing is lost: every sentence is in some chunk.
    joined = " ".join(chunks)
    assert all(f"number {i} " in joined for i in range(120))


def test_a_short_tail_is_folded_into_the_chunk_before():
    text = " ".join(["Word " * 20 + "end."] * 11) + " Tiny tail here."
    chunks = L.chunk_text(text)
    assert chunks[-1].endswith("Tiny tail here.")
    assert all(words(c) > 10 for c in chunks)


def test_an_overlong_sentence_is_cut_by_words():
    chunks = L.chunk_text(" ".join(["x"] * 900))
    assert all(words(c) <= L.MAX_WORDS for c in chunks)
    assert sum(words(c) for c in chunks) >= 900


def test_tokens_keep_flip_flop_letters_and_undo_contractions():
    assert "d" in L.tokens("A D flip-flop")
    assert "t" not in L.tokens("don't")
    assert L.tokens("completed")[0] == L.tokens("complete")[0]
    assert L.query_terms("What is a T flip-flop?") == ["t", "flip", "flop"]
    # Hinglish scaffolding goes; the topic stays.
    assert L.query_terms("setup time kya hota hai bhai?") == ["setup", "time"]


# ------------------------------------------------------------------ extraction
def test_pdf_pages_scanned_and_blank(tmp_path):
    import fitz
    p = C.write_pdf(tmp_path / "x.pdf", "Title", ["Page one has words about multiplexers and select lines.", None])
    doc = fitz.open(str(p))
    doc.new_page()                                       # a blank page, no image
    doc.saveIncr()
    doc.close()
    ex = L.extract(p)
    assert ex.pages == 3 and ex.scanned == [2] and ex.blank == [3]
    assert [s.loc for s in ex.sections] == ["p. 1"]
    assert any("OCR not available" in n for n in ex.notes)


def test_scanned_pages_are_read_when_a_local_ocr_exists(tmp_path, monkeypatch):
    p = C.write_pdf(tmp_path / "scan.pdf", "Scan", [None])
    monkeypatch.setattr(L, "ocr_available", lambda: True)
    monkeypatch.setattr(L, "_ocr_page", lambda page: "Recognised words from the scanned tutorial sheet.")
    ex = L.extract(p)
    assert ex.ocr == [1] and ex.scanned == []
    assert ex.sections[0].text.startswith("Recognised words")


def test_docx_sections_by_heading_with_pages_from_breaks(tmp_path):
    ex = L.extract(C.write_docx(tmp_path / "lab.docx"))
    locs = [s.loc for s in ex.sections]
    assert ex.title == C.LAB_TITLE and ex.pages == 2
    assert locs[0].startswith("§ Experiment 1: Verification of logic gates") and locs[0].endswith("p. 1")
    assert any(l.startswith("§ Experiment 3") and l.endswith("p. 2") for l in locs)
    viva = [s for s in ex.sections if s.heading == "Viva questions"][0]
    assert "Why is the NAND gate called universal? | Any Boolean function" in viva.text


def test_docx_without_breaks_cites_headings_only(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Only section", level=1)
    d.add_paragraph("Some words about counters.")
    d.save(str(tmp_path / "plain.docx"))
    ex = L.extract(tmp_path / "plain.docx")
    assert ex.pages == 0 and ex.sections[0].page == 0 and ex.sections[0].loc == "§ Only section"


def test_markdown_headings_front_matter_code_and_links(tmp_path):
    p = tmp_path / "notes.md"
    p.write_text(C.NOTES_MD, encoding="utf-8")
    ex = L.extract(p)
    assert ex.title == "Digital Electronics - exam revision"
    locs = [s.loc for s in ex.sections]
    assert "§ FSM design › Moore vs Mealy" in locs and "§ Memories › SRAM vs DRAM" in locs
    assert not any("not a heading" in l for l in locs)            # inside a code fence
    ttl = [s for s in ex.sections if s.heading.endswith("TTL vs CMOS")][0]
    assert "sequential notes" in ttl.text and "[[" not in ttl.text and "**" not in ex.sections[0].text


def test_txt_cites_line_ranges(tmp_path):
    p = tmp_path / "w.txt"
    p.write_text(C.WARRANTY_TXT, encoding="utf-8")
    ex = L.extract(p)
    assert ex.sections[0].loc.startswith("lines 1")
    assert "13 March 2035" in ex.sections[0].text


def test_pptx_slides_with_titles_and_speaker_notes(tmp_path):
    ex = L.extract(C.write_pptx(tmp_path / "s.pptx"))
    assert ex.pages == len(C.SLIDES)
    third = [s for s in ex.sections if s.page == 3][0]
    assert third.loc == "slide 3" and third.heading == "State encoding"
    assert "One-hot encoding" in third.text and "Speaker notes: Gray encoding" in third.text


def test_unsupported_types_are_refused():
    with pytest.raises(ValueError):
        L.extract("photo.jpg")


# ------------------------------------------------------------------ ingest
@pytest.fixture
def corpus(tmp_path):
    return C.build_corpus(tmp_path / "corpus")


def test_ingest_a_folder_skips_what_is_not_notes(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder)
    rep = lib.ingest(corpus)
    files = {r["file"] for r in rep["results"]}
    assert rep["counts"] == {"added": 9}
    assert "circuit.png" not in files and "~$Digital_Lab_Manual.docx" not in files
    assert "workspace.md" not in files                       # inside .obsidian
    st = lib.status()
    assert st["documents"] == 9 and st["scanned_pages"] == 1 and st["mode"] == "hybrid"


def test_reingesting_an_unchanged_file_is_a_no_op(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder)
    lib.ingest(corpus)
    gen, rows, calls = lib._load()["generation"], lib._table().count_rows(), fake_embedder.calls
    rep = lib.ingest(corpus)
    assert rep["counts"] == {"unchanged": 9}
    assert lib._load()["generation"] == gen and lib._table().count_rows() == rows
    assert fake_embedder.calls == calls                      # nothing re-embedded
    # Touching a file without changing it costs a hash, not an embedding.
    md = corpus / "Exam_Revision_Notes.md"
    os.utime(md, (time.time() + 5, time.time() + 5))
    assert lib.ingest_file(md)["status"] == "unchanged"
    assert fake_embedder.calls == calls


def test_a_changed_file_replaces_its_passages(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder)
    lib.ingest(corpus)
    md = corpus / "Exam_Revision_Notes.md"
    old = next(d for d in lib.documents() if d["file"] == md.name)
    md.write_text("# Fresh\nThe notes now only talk about the Schmitt trigger and hysteresis.\n", encoding="utf-8")
    r = lib.ingest_file(md)
    assert r["status"] == "updated" and r["doc_id"] != old["doc_id"]
    docs = [d for d in lib.documents() if d["file"] == md.name]
    assert len(docs) == 1 and docs[0]["chunks"] == 1
    hits = lib.search("Schmitt trigger hysteresis", 3)
    assert hits[0]["file"] == md.name and "Schmitt" in hits[0]["text"]
    assert not any("SRAM" in h["text"] for h in lib.search("SRAM refresh", 10) if h["file"] == md.name)


def test_duplicates_and_moved_files(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder)
    src = corpus / "home" / "AC_warranty_card.txt"
    assert lib.ingest_file(src)["status"] == "added"
    copy = tmp_path / "copy.txt"
    copy.write_bytes(src.read_bytes())
    assert lib.ingest_file(copy)["status"] == "duplicate"
    assert len(lib.documents()) == 1
    moved = tmp_path / "moved_warranty.txt"
    src.rename(moved)
    copy.unlink()
    r = lib.ingest_file(moved)
    assert r["status"] == "moved"
    d = lib.documents()[0]
    assert d["file"] == "moved_warranty.txt" and Path(d["path"]) == moved.resolve()
    assert lib.search("compressor warranty", 1)[0]["file"] == "moved_warranty.txt"


def test_editing_one_of_two_identical_files_keeps_the_other(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder)
    a = corpus / "home" / "AC_warranty_card.txt"
    b = tmp_path / "copy_of_card.txt"
    b.write_bytes(a.read_bytes())
    lib.ingest_file(a)
    assert lib.ingest_file(b)["status"] == "duplicate"
    # Edit the ORIGINAL: the copy still holds the old text, so it keeps the passages.
    a.write_text("Fridge warranty: two years on the cooling unit.\n", encoding="utf-8")
    assert lib.ingest_file(a)["status"] == "updated"
    files = {d["file"] for d in lib.documents()}
    assert files == {"AC_warranty_card.txt", "copy_of_card.txt"}
    assert lib.search("compressor warranty 2035", 1)[0]["file"] == "copy_of_card.txt"
    assert lib.search("fridge cooling unit", 1)[0]["file"] == "AC_warranty_card.txt"


def test_remove_forgets_a_document(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder)
    lib.ingest(corpus)
    doc = next(d for d in lib.documents() if d["file"] == "AC_warranty_card.txt")
    assert lib.remove(doc["doc_id"]) is True
    assert lib.remove(doc["doc_id"]) is False
    assert lib.remove("'; drop table --") is False
    assert all(h["file"] != "AC_warranty_card.txt" for h in lib.search("compressor warranty AC", 10))


def test_without_an_embedder_it_runs_on_keywords_and_catches_up_later(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, None)
    lib.ingest(corpus)
    st = lib.status()
    assert st["mode"] == "keyword-only" and st["embedder"]["ready"] is False
    r = lib.retrieve("What does the 74LS138 do?")
    assert r.mode == "keyword-only" and r.best_distance is None
    assert r.hits[0]["file"] == "DE_L3_Combinational.pdf" and r.hits[0]["page"] == 2
    # Later the embedder is there: the stored zero vectors are replaced.
    lib2 = make_lib(tmp_path, fake_embedder)
    assert lib2.sync_vectors() == 9
    assert all(d["embedder"] == fake_embedder.name for d in lib2.documents())
    assert lib2.search("74LS138 decoder", 5, mode="vector")


def test_an_embedder_that_fails_degrades_to_keywords(tmp_path, corpus):
    class Broken:
        name, dim = "broken", 384

        def encode(self, texts, is_query=False):
            raise OSError("model files missing")
    lib = make_lib(tmp_path, Broken())
    lib.ingest(corpus / "home")
    assert lib.status()["mode"] == "keyword-only"
    assert "model files missing" in lib.status()["embedder"]["error"]
    assert lib.retrieve("compressor warranty AC").hits


def test_a_new_embedder_dimension_rebuilds_the_table(tmp_path, corpus, fake_embedder):
    class Small(type(fake_embedder)):
        name, dim = "test-hashing-64", 64

        def encode(self, texts, is_query=False):
            return super().encode(texts, is_query)[:, :64]
    lib = make_lib(tmp_path, fake_embedder)
    lib.ingest(corpus / "home")
    lib2 = make_lib(tmp_path, Small())
    assert lib2.sync_vectors() == 2
    assert lib2._table_dim(lib2._table()) == 64
    assert lib2.search("compressor warranty", 1, mode="vector")[0]["file"] == "AC_warranty_card.txt"


def test_ingest_reports_progress_and_refuses_missing_paths(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder)
    seen = []
    lib.ingest(corpus / "home", progress=seen.append)
    assert seen[0]["done"] == 0 and seen[-1]["done"] == seen[-1]["total"] == 2
    with pytest.raises(FileNotFoundError):
        lib.ingest(tmp_path / "nope")


# ------------------------------------------------------------------ asking
def test_an_empty_library_abstains(tmp_path, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder)
    r = lib.retrieve("What is a flip-flop?")
    assert r.abstained and r.reason == "the library is empty"


class RecordingLLM:
    model_name = "recording fake"

    def __init__(self, reply):
        self.reply, self.calls = reply, []

    def __call__(self, system, user):
        self.calls.append((system, user))
        return self.reply(user) if callable(self.reply) else self.reply


def test_ask_builds_a_numbered_prompt_and_maps_citations_to_pages(tmp_path, corpus, fake_embedder):
    llm = RecordingLLM("It is a 3-to-8 decoder with active-low outputs [1]. It is used for address decoding [1][9].")
    lib = make_lib(tmp_path, fake_embedder, llm)
    lib.ingest(corpus)
    res = lib.ask("What does the 74LS138 do?")
    system, user = llm.calls[0]
    assert "[1] DE_L3_Combinational.pdf, p. 2" in user and user.rstrip().endswith("QUESTION: What does the 74LS138 do?")
    assert L.NOT_IN_NOTES in system and "square brackets" in system
    assert res["used"] == [1] and res["invalid"] == [9] and not res["abstained"]
    src = res["sources"][0]
    assert src["n"] == 1 and src["file"] == "DE_L3_Combinational.pdf" and src["page"] == 2
    assert src["path"].endswith("DE_L3_Combinational.pdf") and "74LS138" in src["text"]
    assert res["llm"] == "recording fake"


def test_the_gate_never_calls_the_model(tmp_path, corpus, fake_embedder):
    llm = RecordingLLM("Ethanol boils at 78 degrees [1].")
    lib = make_lib(tmp_path, None, llm)                     # keyword-only: coverage decides
    lib.ingest(corpus)
    res = lib.ask("What is the boiling point of ethanol?")
    assert res["abstained"] and res["answer"] == L.NOT_IN_NOTES and res["used"] == []
    assert llm.calls == []


def test_the_model_saying_it_does_not_know_is_an_abstention(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder, RecordingLLM(L.NOT_IN_NOTES))
    lib.ingest(corpus)
    res = lib.ask("What is the propagation delay of the 74LS138 in nanoseconds?")
    assert res["abstained"] and not res["uncited"]


def test_an_uncited_answer_is_flagged(tmp_path, corpus, fake_embedder):
    lib = make_lib(tmp_path, fake_embedder, RecordingLLM("A decoder, obviously."))
    lib.ingest(corpus)
    res = lib.ask("What does the 74LS138 do?")
    assert res["uncited"] and res["used"] == []


def test_a_model_that_is_down_still_returns_the_passages(tmp_path, corpus, fake_embedder):
    def down(system, user):
        raise RuntimeError("local model unreachable (ConnectError)")
    lib = make_lib(tmp_path, fake_embedder, down)
    lib.ingest(corpus)
    res = lib.ask("What does the 74LS138 do?")
    assert res["answer"] is None and "unreachable" in res["llm_error"]
    assert res["sources"][0]["file"] == "DE_L3_Combinational.pdf"


@pytest.mark.parametrize("text,expected", [
    ("A [1]. B [2][3].", [1, 2, 3]),
    ("A [1, 3] and [2-4].", [1, 3, 2, 4]),
    ("A [2–3] then [1 and 5].", [2, 3, 1, 5]),
    ("No citations at all.", []),
    ("Array a[i] is not a citation [x].", []),
])
def test_cited_numbers(text, expected):
    assert L.cited_numbers(text) == expected


def test_parse_answer_flags():
    assert L.parse_answer(L.NOT_IN_NOTES, 5)["abstained"]
    assert L.parse_answer("I do not know.", 5)["abstained"]
    p = L.parse_answer("The notes don't cover the delay, but the chip decodes [2].", 5)
    assert not p["abstained"] and p["used"] == [2]
    assert L.parse_answer("Something [7].", 5) == {"answer": "Something [7].", "used": [], "invalid": [7],
                                                   "abstained": False, "uncited": True, "caveat": ""}


def test_a_cited_answer_that_also_refuses_keeps_the_answer_and_says_so():
    """Seen live: both definitions, cited, then the refusal sentence verbatim."""
    raw = ("Setup time is before the clock edge [1]. Hold time is after it [1]. "
           "I don't know - your notes don't cover that.")
    p = L.parse_answer(raw, 5)
    assert p["answer"] == "Setup time is before the clock edge [1]. Hold time is after it [1]."
    assert p["used"] == [1] and not p["abstained"] and p["caveat"] == L.CAVEAT_NOT_FULLY
    # The refusal on its own is still a refusal, with no caveat.
    q = L.parse_answer(L.NOT_IN_NOTES, 5)
    assert q["abstained"] and q["caveat"] == "" and q["answer"] == L.NOT_IN_NOTES


def test_speakable_names_the_sources_instead_of_reading_brackets():
    res = {"answer": "Pin 7 is ground [1]. Pin 14 is Vcc [1][2].", "used": [1, 2], "abstained": False,
           "sources": [{"n": 1, "file": "Digital_Lab_Manual.docx", "loc": "§ Experiment 1: Verification of logic gates · p. 1"},
                       {"n": 2, "file": "DE_L2_Boolean_Kmaps.pdf", "loc": "p. 4"}]}
    s = L.speakable(res)
    assert "[" not in s and s.startswith("Pin 7 is ground. Pin 14 is Vcc.")
    assert "From your notes: Digital Lab Manual, Experiment 1: Verification of logic gates, page 1; DE L2 Boolean Kmaps, page 4." in s
    assert L.speakable({"answer": L.NOT_IN_NOTES, "abstained": True}) == L.NOT_IN_NOTES


def test_context_block_is_empty_for_unrelated_questions(tmp_path, corpus):
    lib = make_lib(tmp_path, None)
    lib.ingest(corpus)
    assert L.context_block(lib, "Who won the 2011 cricket world cup final?") == ""
    block = L.context_block(lib, "Which pin of the 7400 chip is ground?")
    assert block.startswith("FROM HIS STUDY NOTES") and "Digital_Lab_Manual.docx" in block


def test_default_root_follows_the_environment(monkeypatch, tmp_path):
    monkeypatch.setenv("JARVIS_LIBRARY_DIR", str(tmp_path / "elsewhere"))
    assert L.default_root() == tmp_path / "elsewhere"
    monkeypatch.delenv("JARVIS_LIBRARY_DIR")
    assert L.default_root().parts[-3:] == ("data", "yaadash", "library")


def test_a_second_handle_sees_new_documents(tmp_path, corpus, fake_embedder):
    """The app ingests; another process (the voice brain) must see the result
    without a restart: the manifest is re-read when its file changes."""
    a = make_lib(tmp_path, fake_embedder)
    b = make_lib(tmp_path, fake_embedder)
    assert b.retrieve("compressor warranty AC").reason == "the library is empty"
    a.ingest(corpus / "home")
    time.sleep(0.05)
    assert b.retrieve("compressor warranty AC").hits[0]["file"] == "AC_warranty_card.txt"
